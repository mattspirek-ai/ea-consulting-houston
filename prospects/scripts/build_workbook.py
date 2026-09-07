#!/usr/bin/env python3
"""
build_workbook.py — assemble prospects/prospects.xlsx from prospects/verify/*.json

Sheets:
  README            methodology, verification gates, counts, pull date
  Verified          rows passing all gates (the outreach list)
  Partially Verified rows missing >=1 gate, with reasons (do not pitch until cleared)
  Rejected          rows failing a fatal gate, duplicates, or researcher-rejected
  Sources           one row per (company, field, source URL) — full audit trail
  Segments          counts by vertical, metro, trigger, tier, email kind

Also writes prospects/prospects.json (merged, with verification block) and
prospects/prospects.xlsx.b64 (base64 sidecar for the text-only GitHub integration).
"""
import base64
import json
import os
import time
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
V = os.path.join(ROOT, "verify")
PULL = time.strftime("%Y-%m-%d")

COLS = [
    ("company_name", "Company"), ("decision_maker_name", "Decision-maker"), ("decision_maker_title", "Title"),
    ("vertical", "Vertical"), ("sub_vertical", "Sub-vertical"), ("metro", "Metro"), ("headquarters_city", "City"),
    ("address", "Address"), ("employee_count_estimate", "Employees (est.)"), ("revenue_estimate", "Revenue (est.)"),
    ("phone", "Phone (verified)"), ("contact_email", "Named email (verified)"), ("company_email", "Company email (verified)"),
    ("linkedin_url", "LinkedIn"), ("website", "Website"),
    ("trigger_signal", "Trigger"), ("trigger_detail", "Trigger detail"), ("trigger_date", "Trigger date"), ("trigger_source", "Trigger source"),
    ("suggested_tier", "Suggested tier"), ("icp_fit_rationale", "ICP fit rationale"),
    ("confidence_score", "Confidence (0-100)"), ("status", "Status"), ("_reasons", "Verification notes"),
    ("_email_kind", "Email kind"), ("_segment", "Research segment"), ("notes", "Researcher notes"),
]
SOURCE_FIELDS = ["address_source", "employee_count_source", "revenue_source", "contact_email_source",
                 "company_email_source", "phone_source", "trigger_source"]

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {"Verified": "E3F4E1", "Partially Verified": "FFF4D6", "Rejected": "F9E0E0"}


def load(name):
    p = os.path.join(V, name)
    return json.load(open(p)) if os.path.exists(p) else []


def flat(r):
    d = dict(r)
    d["_reasons"] = " | ".join(r.get("reasons") or [])
    d["_email_kind"] = (r.get("verification") or {}).get("evidence", {}).get("email_kind")
    return d


def write_table(ws, rows, cols):
    ws.append([h for _, h in cols])
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="top")
    for r in rows:
        f = flat(r)
        ws.append([("" if f.get(k) is None else (", ".join(f[k]) if isinstance(f[k], list) else f[k])) for k, _ in cols])
        fill = STATUS_FILL.get(f.get("status"))
        if fill:
            ws.cell(row=ws.max_row, column=[k for k, _ in cols].index("status") + 1).fill = PatternFill("solid", fgColor=fill)
    widths = {"Company": 30, "Decision-maker": 22, "Title": 26, "Address": 34, "Trigger detail": 48,
              "ICP fit rationale": 48, "Verification notes": 48, "Website": 30, "LinkedIn": 34, "Trigger source": 34}
    for i, (_, h) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 16)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def main():
    verified, partial, rejected = load("verified.json"), load("partial.json"), load("rejected.json")
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    readme = [
        ["EA Consulting Houston — Verified Prospect List"],
        [f"Pull / verification date: {PULL}"],
        [""],
        ["Counts", f"Verified: {len(verified)}", f"Partially Verified: {len(partial)}", f"Rejected: {len(rejected)}"],
        [""],
        ["Verification gates (a row is Verified only when ALL pass):"],
        ["G1", "Website live: final HTTP status < 400 after redirects; bot-blocked sites go to browser spot-check, never auto-pass"],
        ["G2", "Decision-maker surname found on at least one fetched official (non-LinkedIn) page"],
        ["G3", "Two title sources on distinct domains, each fetchable or a well-formed LinkedIn profile URL"],
        ["G4", "Phone: 10-digit number found, in formatted form, on the company website or the cited source page"],
        ["G5", "Email: published on the cited page or company site AND domain has MX records (DNS-over-HTTPS) AND not free-mail"],
        ["G6", "Trigger: source URL reachable and trigger detail present"],
        ["G7", "Metro in scope (Houston first; DFW / Austin / San Antonio expansion)"],
        ["G8", "Size band overlaps 5–150 employees (null allowed and flagged, never estimated)"],
        [""],
        ["Data honesty rules (repo decision D-006): nulls are never filled; every field has a source URL (see Sources sheet);"],
        ["corrections are appended, never silently edited. Contact data is limited to publicly published business channels."],
        ["Mailbox-level SMTP verification is not possible from the build environment; G5 is publication + MX evidence."],
        ["Confidence = share of the six hard gates (G1–G6) passed."],
    ]
    for line in readme:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 120

    for title, rows in (("Verified", verified), ("Partially Verified", partial), ("Rejected", rejected)):
        write_table(wb.create_sheet(title), rows, COLS)

    src = wb.create_sheet("Sources")
    src.append(["Company", "Field", "Value", "Source URL", "Status"])
    for c in src[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in verified + partial + rejected:
        for f in SOURCE_FIELDS:
            if r.get(f):
                src.append([r.get("company_name"), f.replace("_source", ""), str(r.get(f.replace("_source", "")) or ""), r[f], r.get("status")])
        for u in (r.get("title_sources") or []):
            src.append([r.get("company_name"), "title", r.get("decision_maker_title") or "", u, r.get("status")])
    for col, w in zip("ABCDE", (30, 16, 40, 70, 18)):
        src.column_dimensions[col].width = w
    src.auto_filter.ref = src.dimensions

    seg = wb.create_sheet("Segments")
    for label, key in (("Vertical", "vertical"), ("Metro", "metro"), ("Trigger", "trigger_signal"),
                       ("Suggested tier", "suggested_tier"), ("Research segment", "_segment")):
        seg.append([f"Verified by {label}", "Count"])
        seg.cell(row=seg.max_row, column=1).font = Font(bold=True)
        for k, v in Counter(r.get(key) for r in verified).most_common():
            seg.append([k, v])
        seg.append([])
    seg.append(["Verified by email kind", "Count"])
    seg.cell(row=seg.max_row, column=1).font = Font(bold=True)
    for k, v in Counter(flat(r)["_email_kind"] for r in verified).most_common():
        seg.append([k, v])
    seg.column_dimensions["A"].width = 32

    out = os.path.join(ROOT, "prospects.xlsx")
    wb.save(out)
    json.dump({"pulled_at": PULL, "verified": verified, "partially_verified": partial, "rejected": rejected},
              open(os.path.join(ROOT, "prospects.json"), "w"), indent=2)
    with open(out, "rb") as f, open(out + ".b64", "w") as g:
        g.write(base64.b64encode(f.read()).decode())
    print(f"wrote {out} ({os.path.getsize(out)} bytes); Verified {len(verified)} / Partial {len(partial)} / Rejected {len(rejected)}")


if __name__ == "__main__":
    main()
