import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

R = "../research/"
nat = json.load(open(R + "national_competitors.json"))
hou = json.load(open(R + "houston_competitors.json"))
dem = json.load(open(R + "demand_keywords_pricing.json"))

wb = Workbook()
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D0D7E2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def j(v):
    if v is None:
        return "not available"
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (list, dict)):
        if isinstance(v, list):
            if all(isinstance(x, str) for x in v):
                return "; ".join(v)
            return "\n".join(json.dumps(x, ensure_ascii=False) if not isinstance(x, str) else x for x in v)
        return "; ".join(f"{k}: {val}" for k, val in v.items())
    return v

def sheet(title, headers, rows, widths=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="7A4B00")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(headers), 12))
        r0 = 2
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.alignment = WRAP; cell.border = BORDER
    for i, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=j(v))
            cell.alignment = WRAP; cell.border = BORDER
            if isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v; cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    for c in range(1, len(headers) + 1):
        w = (widths[c - 1] if widths and c - 1 < len(widths) else 22)
        ws.column_dimensions[get_column_letter(c)].width = w
    return ws

# ---------- README ----------
ws = wb.active; ws.title = "README"
readme = [
    ("Houston + Virtual Executive Assistant Market Research — Workbook", True),
    ("Prepared 2026-09-07 for a fractional/virtual EA services launch (Houston base, national virtual footprint).", False),
    ("", False),
    ("HOW TO READ THIS WORKBOOK", True),
    ("• Every number that is an estimate is labeled approximate and carries a source URL and pull date.", False),
    ("• 'not available' means the value could not be sourced. Nothing has been guessed or filled with a neutral value.", False),
    ("• Traffic figures come from Exa entity data (approx. June 2026 snapshot). Treat as order-of-magnitude, not precise.", False),
    ("• 'Who ranks now' columns are based on the Exa search index, not a live Google SERP.", False),
    ("", False),
    ("KNOWN DATA GAPS", True),
    ("1. Keyword monthly volumes are all 'not available' — no free public tool returned numeric volume during research. Pull from Google Keyword Planner (free with a Google Ads account) before finalizing SEO priorities.", False),
    ("2. Domain ratings (DR/DA) are 'not available' throughout — pull from Ahrefs/Moz free checkers manually.", False),
    ("3. The '547 employed' Houston EA figure (SalaryMetro, SOC 43-6014) is UNVERIFIED and likely wrong; verify against BLS OEWS Houston May 2024 release.", False),
    ("4. Google Maps pack for Houston queries was not captured (Places API unavailable during research).", False),
    ("5. Priority VA, Assistantly, Great Assistant were not retrievable; Equivity returned 404 (may be closed).", False),
    ("", False),
    ("SHEETS", True),
    ("Competitors — 38 rows: national virtual EA firms + Houston agencies, solos, and staffing firms, one schema.", False),
    ("Keywords — 61 terms across 7 clusters with intent, who ranks now, content type needed, and priority for a new entrant.", False),
    ("Pricing — 13 sourced benchmark data points + in-house cost derivation.", False),
    ("Pricing Recommendation — derived tier options for the new Houston entrant (recommendations, not sourced facts).", False),
    ("Buyer Signals — pain points, praise, objections, guarantees seen (with sources).", False),
    ("Houston Demand — sourced market context stats for Houston (GHP, Dallas Fed, BLS).", False),
    ("Houston SERP Snapshot — who appears for key Houston queries (search-index based).", False),
    ("Gaps & Positioning — market gaps nobody fills and the recommended positioning.", False),
    ("Sources — every URL referenced, deduplicated.", False),
]
for i, (t, bold) in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = Font(bold=bold, size=13 if i == 1 else 11); c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 140

# ---------- Competitors ----------
hdr = ["Name", "URL", "Market", "Group", "HQ / Location", "Founded", "Business model", "Target customer", "Service scope",
       "Talent location", "Pricing model", "Pricing public?", "Price low (USD)", "Price high (USD)", "Price notes",
       "Est. monthly visits (approx.)", "Visits source", "Domain rating (approx.)", "DR source", "Google rating", "Google reviews",
       "Other reviews", "Top keywords targeted", "Positioning statement", "Differentiators", "Notable website features",
       "Trust signals", "Website quality (1-5)", "Weaknesses / gaps", "LinkedIn presence", "Notes", "Sources", "Pulled at"]
rows = []
for c in nat["competitors"]:
    rows.append([c["name"], c["url"], "National / Virtual", "National", c.get("hq_location"), c.get("founded_year"), c.get("business_model"),
                 c.get("target_customer"), c.get("service_scope"), c.get("talent_location"), c.get("pricing_model"), c.get("pricing_public"),
                 c.get("price_low_usd"), c.get("price_high_usd"), c.get("price_notes"), c.get("est_monthly_visits"), c.get("est_visits_source"),
                 c.get("domain_rating"), c.get("dr_source"), None, None, c.get("review_platforms_and_counts"), c.get("top_keywords"),
                 c.get("positioning_statement"), c.get("differentiators"), c.get("notable_website_features"), c.get("trust_signals"), None,
                 c.get("weaknesses_or_gaps"), None, c.get("notes"), c.get("sources"), nat["pulled_at"]])
# Delegate Solutions appears in pricing + SERP data but not in the national list — add from sourced pricing data
rows.append(["Delegate Solutions", "https://www.delegatesolutions.com", "National / Virtual", "National", "US (remote)", None,
             "agency-team (3-person team model)", "Founders, executives", "Fractional EA (team model)", "US-based employees",
             "monthly retainer (tiered hours)", True, 2299, 3959, "$2,299–$3,959/mo for 25–45 hrs ($88–$92/hr effective); $499 one-time onboarding",
             None, None, None, None, 4.9, "50+", None, ["fractional executive assistant", "executive assistant pricing", "fractional EA cost"],
             "Fractional executive assistant team for founders", "Transparent pricing; team model = no single point of failure", "Public pricing page",
             "4.9 Google rating (50+ reviews)", None, "Highest effective hourly in US segment", None,
             "Added from pricing research; ranks #1 in index for 'fractional executive assistant'",
             ["https://www.delegatesolutions.com/fractional-executive-assistant-pricing"], dem["pulled_at"][:10]])
GROUP = {"A": "A — Houston agency / Houston-targeted", "B": "B — Houston solo operator", "C": "C — Houston staffing / placement"}
for c in hou["competitors"]:
    rows.append([c["name"], c["url"], "Houston", GROUP.get(c.get("group"), c.get("group")), c.get("location_specific"), None, c.get("business_model"),
                 c.get("target_customer"), c.get("service_scope"), c.get("talent_location"), c.get("pricing_model"), c.get("pricing_public"),
                 c.get("price_low_usd"), c.get("price_high_usd"), c.get("price_notes"), c.get("est_monthly_visits"), c.get("est_visits_source"),
                 c.get("domain_rating"), c.get("dr_source"), c.get("google_rating"), c.get("google_review_count"), c.get("other_reviews"),
                 c.get("top_keywords"), c.get("positioning_statement"), c.get("differentiators"), c.get("notable_website_features"),
                 c.get("trust_signals"), (f"{c['website_quality_score_1to5']} — {c.get('website_quality_justification','')}" if c.get("website_quality_score_1to5") else None),
                 c.get("weaknesses_or_gaps"), c.get("linkedin_presence"), c.get("notes"), c.get("sources"), c.get("pulled_at")])
sheet("Competitors", hdr, rows,
      widths=[26, 34, 16, 26, 28, 9, 30, 30, 30, 26, 26, 10, 12, 12, 40, 14, 34, 12, 20, 10, 10, 30, 40, 44, 44, 44, 40, 34, 44, 24, 44, 50, 12],
      note="Traffic and DR values are APPROXIMATE (Exa entity data ~Jun 2026). 'not available' = could not be sourced; nothing is estimated.")

# ---------- Keywords ----------
hdr = ["Keyword", "Cluster", "Intent", "Est. monthly US volume (approx.)", "Volume source", "Difficulty note", "Who ranks now (search-index based)",
       "Content type needed", "Priority for new entrant (1-5)", "Priority rationale"]
rows = [[k["term"], k["cluster"], k["intent"], k["est_monthly_volume_us"], k["volume_source"], k["difficulty_note"], k["who_ranks_now"],
         k["content_type_needed"], k["priority_1to5"], k["priority_rationale"]] for k in
        sorted(dem["keywords"], key=lambda x: (-x["priority_1to5"], x["cluster"], x["term"]))]
sheet("Keywords", hdr, rows, widths=[38, 16, 14, 16, 30, 40, 44, 20, 12, 60],
      note="Volumes are 'not available' for every term — no free tool returned numbers. Pull from Google Keyword Planner before locking SEO priorities. 'Who ranks now' is Exa-index based, not a live SERP.")

# ---------- Pricing ----------
hdr = ["Provider / source", "Model", "Price", "Unit", "Hours included", "Effective $/hr", "Talent location", "Source URL", "Date seen", "Notes"]
rows = [[p.get("provider_or_source"), p.get("model"), p.get("price"), p.get("price_unit"), p.get("hours_included"), p.get("per_hour_effective"),
         p.get("talent_location"), p.get("url"), p.get("date_seen"), p.get("notes")] for p in dem["pricing_benchmarks"]]
ls = dem["labor_stats"]
rows.append(["BLS OES 43-6011 (national)", "Executive Secretaries & Exec Admin Assistants — median annual wage", ls["national_bls_oes_43_6011"]["median_annual_wage_p50"], "per_year",
             None, ls["national_bls_oes_43_6011"]["median_hourly_wage_p50"], "US in-house W-2", ls["national_bls_oes_43_6011"]["source"], "May 2023 data",
             f"Mean ${ls['national_bls_oes_43_6011']['mean_annual_wage']}; P25 ${ls['national_bls_oes_43_6011']['p25_annual']}; P75 ${ls['national_bls_oes_43_6011']['p75_annual']}; P90 ${ls['national_bls_oes_43_6011']['p90_annual']}"])
rows.append(["Indeed — Houston EA salaries", "Average base salary, Houston TX", ls["houston_ea_salary_indeed"]["average_base_salary"], "per_year", None, None,
             "Houston in-house W-2", ls["houston_ea_salary_indeed"]["source"], ls["houston_ea_salary_indeed"]["survey_date_updated"],
             f"Range ${ls['houston_ea_salary_indeed']['low']}–${ls['houston_ea_salary_indeed']['high']}; n={ls['houston_ea_salary_indeed']['n_salaries']}"])
rows.append(["Robert Half 2026 — Houston EA", "Salary range low/mid/high", ls["houston_ea_salary_robert_half"]["range_mid"], "per_year", None, None,
             "Houston in-house W-2", ls["houston_ea_salary_robert_half"]["source"], "2026",
             f"Low ${ls['houston_ea_salary_robert_half']['range_low']}; High ${ls['houston_ea_salary_robert_half']['range_high']}"])
rows.append(["SalaryMetro citing BLS OEWS 2024 — Houston MSA", "Median annual wage (UNVERIFIED — SOC 43-6014 used; '547 employed' figure looks wrong)", ls["houston_msa_ea_salary"]["median_annual_wage"],
             "per_year", None, ls["houston_msa_ea_salary"]["median_hourly_wage"], "Houston in-house W-2", ls["houston_msa_ea_salary"]["source"], "2024 data",
             "VERIFY against BLS: " + ls["houston_msa_ea_salary"]["source_note"]])
fl = ls["fully_loaded_cost_in_house_houston"]
rows.append(["DERIVED — fully loaded in-house Houston EA", "Recommendation/derivation, not a sourced figure", fl["estimated_fully_loaded_monthly"], "per_month", 173,
             round(fl["estimated_fully_loaded_monthly"] / 173, 2), "Houston in-house W-2", None, "2026-09-07", fl["derivation_basis"]])
sheet("Pricing", hdr, rows, widths=[34, 44, 10, 10, 10, 10, 26, 44, 12, 70],
      note="All rows sourced except those labeled DERIVED. Effective $/hr computed only where hours were published.")

# ---------- Pricing Recommendation ----------
pr = dem["pricing_recommendations"]
hdr = ["Option", "Tier", "Hours / month", "Monthly price (USD)", "Effective $/hr", "Rationale", "Competitive gap", "Risks"]
rows = []
for key in ["option_a_premium_boutique", "option_b_mid_market"]:
    o = pr[key]
    for t in ["entry_tier", "growth_tier", "full_tier"]:
        rows.append([o["label"], o[t]["name"], o[t]["hours_per_month"], o[t]["monthly"], o[t]["effective_per_hour"], o["rationale"], o.get("competitive_gap"), o["risks"]])
rows.append(["RECOMMENDED APPROACH", None, None, None, None, pr["recommended_approach"], None, None])
rows.append(["Competitor summary", None, None, None, None, pr["competitor_pricing_summary"], None, None])
sheet("Pricing Recommendation", hdr, rows, widths=[30, 18, 12, 14, 12, 70, 50, 50],
      note=pr["label"])

# ---------- Buyer Signals ----------
b = dem["buyer_signals"]
hdr = ["Type", "Provider", "Theme", "Signal", "Source URL"]
rows = []
for x in b["pain_points"]:
    rows.append(["Pain point", x.get("source_provider"), x.get("theme"), x.get("pain_point"), x.get("source_url")])
for x in b["praise"]:
    rows.append(["Praise", x.get("source_provider"), None, x.get("praise_point"), x.get("source_url")])
for x in b["objections"]:
    rows.append(["Objection", "Synthesized across reviews", None, x, None])
for x in b["guarantees_seen"]:
    rows.append(["Guarantee seen", x.get("provider"), None, x.get("guarantee"), x.get("source")])
sheet("Buyer Signals", hdr, rows, widths=[14, 24, 20, 90, 50])

# ---------- Houston Demand ----------
hdr = ["Topic", "Stat", "Source", "URL"]
rows = [[d["topic"], d["stat"], d["source"], d["url"]] for d in hou["demand_context"]]
t = dem["trends_and_demand"]["market_size"]
rows.append(["Global VA services market", t["stat"], f"{t['publisher']} ({t['year_published']})", t["source"]])
sheet("Houston Demand", hdr, rows, widths=[30, 90, 40, 50])

# ---------- Houston SERP Snapshot ----------
hdr = ["Query", "Search type", "Top results (in order)", "Note"]
rows = []
for s in hou["serp_snapshot"]:
    names = [r if isinstance(r, str) else (r.get("name") or r.get("domain") or r.get("url")) for r in s["top_results"]]
    rows.append([s["query"], s.get("search_type"), "\n".join(f"{i+1}. {n}" for i, n in enumerate(names)), s.get("note")])
sheet("Houston SERP Snapshot", hdr, rows, widths=[44, 14, 60, 50],
      note="Search-index based (Exa), not a live Google SERP or Maps pack.")

# ---------- Gaps & Positioning ----------
hdr = ["Category", "Item", "Implication for new Houston entrant"]
rows = []
for g in nat["market_patterns"]["gaps_nobody_fills"]:
    rows.append(["Market gap (national)", g, None])
for p in nat["market_patterns"]["common_promises"]:
    rows.append(["Table-stakes promise", p, "Must match; not a differentiator"])
for k, v in nat["market_patterns"]["pricing_bands"].items():
    rows.append(["Pricing band", f"{k}: {v}", None])
for p in nat["market_patterns"]["common_page_structures"]:
    rows.append(["Common page structure", p, "Mirror the structure; differentiate the content"])
rows.append(["Houston observation", hou["market_observations"], None])
rows.append(["RECOMMENDED POSITIONING", "Houston-rooted, US-based senior fractional EAs for energy, healthcare/TMC, legal, and PE/family-office leaders. Published pricing. Month-to-month. No buyout fees. Matched in 5 business days. Optional Houston in-person hybrid days.",
             "Occupies the only unclaimed intersection: local credibility × US talent × transparent pricing × industry fluency."])
sheet("Gaps & Positioning", hdr, rows, widths=[28, 110, 60])

# ---------- Sources ----------
urls = set()
def harvest(o):
    if isinstance(o, str) and o.startswith("http"): urls.add(o)
    elif isinstance(o, list): [harvest(x) for x in o]
    elif isinstance(o, dict): [harvest(v) for v in o.values()]
harvest(nat); harvest(hou); harvest(dem)
sheet("Sources", ["URL"], [[u] for u in sorted(urls)], widths=[120])

out = "../data/EA_Houston_Virtual_Market_Research.xlsx"
wb.save(out)
print("saved", out, "competitor rows:", len(nat["competitors"]) + 1 + len(hou["competitors"]), "sources:", len(urls))
